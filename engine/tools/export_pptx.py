"""PowerPoint backup of a deck: one full-bleed picture per slide (exactly what the HTML deck shows, animations at their
final state) with the speaker notes in PowerPoint's notes pane.
  .aura/venv/Scripts/python.exe .aura/engine/tools/export_pptx.py "4 - Your slides/<Title>.html" ["<out.pptx>"]"""
import io, json, shutil, subprocess, sys
from pathlib import Path

# A console app launched from a windowless parent opens its OWN console window on Windows.
# Lumi runs this during an export, so without this flag a black box flashes over the app.
NO_WINDOW = getattr(subprocess, 'CREATE_NO_WINDOW', 0)

TOOLS = Path(__file__).resolve().parent
try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass


def aura_root(p: Path):
    for d in [p] + list(p.parents):
        if (d / '.aura').is_dir():
            return d
    return None


def render(deck: Path):
    """Shoot every slide with Edge (shoot_slides.js). Returns (folder, slides.json data)."""
    root = aura_root(deck)
    out = (root / '.aura' / 'temp' / 'export' / deck.stem) if root else deck.parent / ('_' + deck.stem + '_slides')
    node = shutil.which('node')
    if not node:
        raise SystemExit('Node.js is missing, so the slides cannot be rendered. Run "Update Lumi".')
    r = subprocess.run([node, str(TOOLS / 'shoot_slides.js'), str(deck), str(out)], capture_output=True, text=True,
                       encoding='utf-8', errors='replace', stdin=subprocess.DEVNULL, creationflags=NO_WINDOW)
    if r.returncode != 0:
        raise SystemExit((r.stderr or r.stdout).strip() or 'Rendering the slides failed.')
    return out, json.loads((out / 'slides.json').read_text(encoding='utf-8'))


def compact(png: Path):
    """JPEG at high quality is a fraction of the PNG size and looks the same on a projector."""
    from PIL import Image
    buf = io.BytesIO()
    with Image.open(png) as im:
        im.convert('RGB').save(buf, 'JPEG', quality=92, optimize=True, progressive=True, subsampling=0)
    if buf.tell() >= png.stat().st_size:
        return io.BytesIO(png.read_bytes())
    buf.seek(0)
    return buf


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    deck = Path(sys.argv[1]).resolve()
    if deck.is_dir():
        deck = deck / 'index.html'
    if not deck.is_file():
        print(f'Deck not found: {deck}')
        return 2
    out = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else deck.with_suffix('.pptx')
    folder, meta = render(deck)

    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)       # 13.333 x 7.5 in, 16:9
    blank = prs.slide_layouts[6]
    notes_count = 0
    for s in meta['slides']:
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(compact(folder / s['image']), 0, 0, width=prs.slide_width, height=prs.slide_height)
        pic.name = f"Slide {s['number']}"
        pic._element.nvPicPr.cNvPr.set('descr', s.get('title') or f"Slide {s['number']}")
        if s.get('notes'):
            slide.notes_slide.notes_text_frame.text = s['notes']
            notes_count += 1
    prs.core_properties.title = meta.get('title') or deck.stem
    prs.core_properties.author = 'Lumi'
    tmp = out.with_name(out.name + '.part')
    prs.save(tmp)
    tmp.replace(out)

    check = Presentation(out)
    n = len(check.slides)
    root = aura_root(deck)
    shown = str(out.relative_to(root) if root and str(out).startswith(str(root)) else out).replace('\\', '/')
    print(f'PowerPoint saved: {shown}')
    print(f'  {n} slides (pictures of the deck), speaker notes on {notes_count}, {out.stat().st_size / 1048576:.2f} MB')
    if n != len(meta['slides']):
        print('  Slide count does not match the deck.')
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
