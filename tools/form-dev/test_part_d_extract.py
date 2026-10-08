"""Part D: every upload kind ends with text to read or a plain warning, scanned pages included.
Builds its own fixtures (a scanned PDF drawn from text, a Word file with a pasted table picture, a .potx, an .svg, a broken
.heic) in a scratch folder, runs extract_text.py on them and checks the manifest. Windows only for the machine-reading checks.
  python tools/form-dev/test_part_d_extract.py [scratch folder]      (default: X:\\aura-dev-d\\t_extract)
Use the private venv's python (needs the packages in setup/aura.config.json)."""
import io, json, os, shutil, subprocess, sys, zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
TOOL = HERE.parent.parent / 'engine' / 'tools' / 'extract_text.py'
root = Path(sys.argv[1] if len(sys.argv) > 1 else r'X:\aura-dev-d\t_extract')
bad = 0


def check(name, ok, info=None):
    global bad
    bad += not ok
    print(('PASS ' if ok else 'FAIL ') + name + ('' if ok else f'   <- {info!r}'[:400]))


sys.path.insert(0, str(TOOL.parent))
import extract_text as X                                                    # noqa: E402

# ---- the watermark that hid the owner's scanned report: unit level, no OCR needed
stamp = ['CamScanner'] * 12
check('a stamp printed on every page does not count as text: all 12 pages are scanned', X.thin_pages(stamp) == list(range(1, 13)))
mixed = ['CamScanner\n' + 'Real text of a written page. ' * 5] + ['CamScanner'] * 3
check('a page with real writing under the same stamp is not scanned', X.thin_pages(mixed) == [2, 3, 4], X.thin_pages(mixed))
check('a one-page PDF with a text layer is not scanned', X.thin_pages(['Abstract. ' * 20]) == [])

# ---- fixtures
from PIL import Image, ImageDraw, ImageFont                                 # noqa: E402
import docx                                                                 # noqa: E402
from pptx import Presentation                                               # noqa: E402

shutil.rmtree(root, ignore_errors=True)
src, out = root / 'files', root / 'text'
for d in ('Report', 'Images and photos', 'Logo and university template', 'Anything else'):
    (src / d).mkdir(parents=True)
try:
    font = ImageFont.truetype('arial.ttf', 44)
except OSError:
    font = ImageFont.load_default()


def page(lines):
    im = Image.new('RGB', (1700, 2200), 'white')
    dr = ImageDraw.Draw(im)
    for k, t in enumerate(lines):
        dr.text((140, 160 + k * 90), t, fill='black', font=font)
    return im.rotate(0.6, fillcolor='white')                                # a scan is never quite straight


p1 = page(['Interim Progress Report', 'Air enters at uniform velocity and 307.15 K.', 'The coolant Reynolds number is about 1830.'])
p2 = page(['Mesh and discretisation uncertainty', 'Two meshes of 2,589,739 elements were solved.'])
p1.save(src / 'Report' / 'scanned.pdf', save_all=True, append_images=[p2], resolution=200)
table = page(['Minimum free-flow area 40.655 mm2', 'Hydraulic diameter 3.0329 mm']).crop((0, 100, 1700, 420))
table.save(src / 'Images and photos' / 'table.png')
dx = docx.Document()
dx.add_paragraph('Table 1 is pasted below as a picture.')
b = io.BytesIO(); table.save(b, 'PNG'); b.seek(0)
dx.add_picture(b, width=docx.shared.Inches(6))
dx.save(src / 'Report' / 'pasted.docx')
prs = Presentation()
prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = 'University template title'
t = io.BytesIO(); prs.save(t)
with zipfile.ZipFile(t) as a, zipfile.ZipFile(src / 'Logo and university template' / 'uni.potx', 'w') as z:
    for n in a.namelist():
        x = a.read(n)
        z.writestr(n, x.replace(b'presentationml.presentation.main+xml', b'presentationml.template.main+xml') if n == '[Content_Types].xml' else x)
(src / 'Images and photos' / 'd.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg"><text>fin pitch 2.50 mm</text></svg>', encoding='utf-8')
(src / 'Images and photos' / 'phone.heic').write_bytes(b'\0\0\0\x18ftypheic' + b'\0' * 200)
(src / 'Anything else' / 'old.doc').write_bytes(b'\xd0\xcf\x11\xe0' + b'\0' * 200)

r = subprocess.run([sys.executable, str(TOOL), str(src), '--out', str(out)], capture_output=True, text=True, encoding='utf-8')
check('extract_text.py runs clean', r.returncode == 0, r.stderr[-300:])
man = json.loads((out / 'manifest.json').read_text(encoding='utf-8'))['files']
txt = lambda rel: (out / (rel + '.txt')).read_text(encoding='utf-8') if (out / (rel + '.txt')).is_file() else ''

s = man.get('Report/scanned.pdf', {})
check('scanned PDF: both pages listed as scanned', s.get('scanned') == [1, 2], s)
check('scanned PDF: a drawn picture per page, in the manifest', sorted(s.get('pagePictures', {})) == ['1', '2'] and all(Path(v).is_file() for v in s.get('pagePictures', {}).values()), s.get('pagePictures'))
if os.name == 'nt' and s.get('ocr') != 'unavailable':
    body = txt('Report/scanned.pdf')
    check('scanned PDF: the text was machine-read (307.15, 1830, 2,589,739)', all(v in body for v in ('307.15', '1830', '2,589,739')), body[:600])
    check('scanned PDF: each page is marked as machine-read and names its picture', body.count('scanned page, text machine-read') == 2 and 'page-001.jpg' in body)
    check('scanned PDF: the warning says check numbers on the page', any('digits in tables are not' in w for w in s.get('warnings', [])), s.get('warnings'))
    check('Word file with a pasted table picture: the picture was machine-read', '40.655' in txt('Report/pasted.docx'), txt('Report/pasted.docx')[-400:])
    check('a loose picture of a table: machine-read', '40.655' in txt('Images and photos/table.png'))
else:
    check('no text recognition here: the scanned PDF says so plainly', any('no text recognition' in w for w in s.get('warnings', [])), s.get('warnings'))
check('.potx template: read', 'University template title' in txt('Logo and university template/uni.potx'))
check('.svg: its labels are text', '2.50' in txt('Images and photos/d.svg'))
h = man.get('Images and photos/phone.heic', {})
check('a .heic this PC cannot open: a plain warning with the fix', any('HEIC' in w and 'JPG' in w for w in h.get('warnings', [])), h)
o = man.get('Anything else/old.doc', {})
check('old .doc: a plain warning with the fix', any('.docx' in w for w in o.get('warnings', [])), o)
check('every file ends with text or a warning (no silence)', all(e.get('chars') or e.get('warnings') for e in man.values()),
      {k: e for k, e in man.items() if not (e.get('chars') or e.get('warnings'))})

print(f'{bad} FAILED' if bad else 'all passed')
sys.exit(1 if bad else 0)
